"""추천 결과 JSON을 초보자가 읽기 쉬운 Excel 보고서로 만듭니다."""

import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


HEADERS = ["rank", "top3", "keyword", "score", "monthly_pc", "monthly_mobile",
           "monthly_total", "trend_7d", "trend_30d", "blog_total",
           "opportunity_index", "sub_keyword_1", "sub_keyword_2", "sub_keyword_3",
           "recommendation_reason", "selected"]


def add_table(ws, start_row, rows):
    for column, header in enumerate(HEADERS, 1):
        cell = ws.cell(start_row, column, header)
        cell.fill = PatternFill("solid", fgColor="14532D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_index, item in enumerate(rows, start_row + 1):
        for column, header in enumerate(HEADERS, 1):
            value = item.get(header, "")
            if header in {"trend_7d", "trend_30d"} and isinstance(value, (int, float)):
                value /= 100
            cell = ws.cell(row_index, column, value)
            cell.alignment = Alignment(vertical="top", wrap_text=header.startswith("sub_") or header == "recommendation_reason")
        for column in (5, 6, 7, 10):
            ws.cell(row_index, column).number_format = "#,##0"
        for column in (8, 9):
            ws.cell(row_index, column).number_format = "+0.0%;-0.0%;0.0%"
        for column in (4, 11):
            ws.cell(row_index, column).number_format = "0.0"
        if item.get("top3"):
            for column in range(1, 17):
                ws.cell(row_index, column).fill = PatternFill("solid", fgColor="FEF3C7")
    if rows:
        ws.conditional_formatting.add(
            f"D{start_row + 1}:D{start_row + len(rows)}",
            ColorScaleRule(start_type="min", start_color="FEE2E2", end_type="max", end_color="DCFCE7"),
        )


def build(input_path: Path, output_dir: Path) -> Path:
    report = json.loads(input_path.read_text(encoding="utf-8"))
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / f"오늘의_블로그_키워드_{report['generation_date']}.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "오늘의 추천"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:P2")
    ws["A1"] = "오늘의 블로그 키워드 리포트"
    ws["A1"].fill = PatternFill("solid", fgColor="14532D")
    ws["A1"].font = Font(size=20, bold=True, color="FFFFFF")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    summary = [("생성일", report["generation_date"]), ("블로그 분야", report["blog_category"]),
               ("분석 후보 수", report["candidate_count"]), ("최근 14일 제외 수", report["excluded_count"])]
    for idx, (label, value) in enumerate(summary):
        col = idx * 2 + 1
        ws.cell(4, col, label).font = Font(bold=True, color="14532D")
        ws.cell(4, col + 1, value)
    ws["A5"], ws["B5"] = "정상 분석 수", report["normal_analysis_count"]
    ws["A5"].font = Font(bold=True, color="14532D")
    ws["C5"], ws["D5"] = "데이터 안내", report.get("data_note", "")
    ws["C5"].font = Font(bold=True, color="14532D")
    ws.merge_cells("D5:P5")
    ws["D5"].alignment = Alignment(wrap_text=True)
    if report.get("insufficient_data_message"):
        ws.merge_cells("A6:P6")
        ws["A6"] = report["insufficient_data_message"]
        ws["A6"].font = Font(bold=True, color="B91C1C")
    ws.merge_cells("A7:P7"); ws["A7"] = "TOP3 상세"
    ws["A7"].font = Font(size=14, bold=True, color="14532D")
    add_table(ws, 8, report["top3"])
    ws.merge_cells("A13:P13"); ws["A13"] = "TOP30 전체"
    ws["A13"].font = Font(size=14, bold=True, color="14532D")
    add_table(ws, 14, report["top30"])
    widths = [8, 9, 20, 10, 15, 17, 16, 13, 13, 17, 18, 22, 24, 22, 58, 12]
    for column, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A15"
    validation = DataValidation(type="list", formula1='"TRUE"', allow_blank=True)
    ws.add_data_validation(validation)
    last_row = max(15, 14 + len(report["top30"]))
    validation.add(f"P15:P{last_row}")

    history = wb.create_sheet("선택 이력")
    history.sheet_view.showGridLines = False
    history.merge_cells("A1:C2"); history["A1"] = "선택 이력"
    history["A1"].fill = PatternFill("solid", fgColor="14532D")
    history["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    history["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for column, header in enumerate(("keyword", "selected_date", "exclude_until"), 1):
        cell = history.cell(4, column, header); cell.fill = PatternFill("solid", fgColor="14532D")
        cell.font = Font(color="FFFFFF", bold=True); cell.alignment = Alignment(horizontal="center")
    for row_index, item in enumerate(report.get("selection_history", []), 5):
        history.cell(row_index, 1, item["keyword"])
        history.cell(row_index, 2, item["selected_date"])
        history.cell(row_index, 3, item["exclude_until"])
    history.column_dimensions["A"].width = 28
    history.column_dimensions["B"].width = 18
    history.column_dimensions["C"].width = 18
    history.freeze_panes = "A5"
    wb.save(output_path)
    return output_path


if __name__ == "__main__":
    result = build(Path(sys.argv[1]), Path(sys.argv[2]))
    # 다시 열어 실제 Excel 압축 구조와 필수 시트가 유효한지 확인합니다.
    checked = load_workbook(result, data_only=False)
    assert checked.sheetnames == ["오늘의 추천", "선택 이력"]
    print(result)
